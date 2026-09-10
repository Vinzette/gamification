import io

import openpyxl
import pandas as pd

from report import build_excel_report


def test_workbook_has_daily_and_monthly_sheets():
    daily = pd.DataFrame([{"Rep": "A", "Total Coins": 40}])
    monthly = pd.DataFrame([{"Rep": "A", "Month": "2026-07", "Total Coins": 40}])
    workbook = openpyxl.load_workbook(io.BytesIO(build_excel_report(daily, monthly)))
    assert workbook.sheetnames == ["Daily", "Monthly"]


def test_header_row_is_bold_on_both_sheets():
    daily = pd.DataFrame([{"Rep": "A", "Total Coins": 40}])
    monthly = pd.DataFrame([{"Rep": "A", "Month": "2026-07", "Total Coins": 40}])
    workbook = openpyxl.load_workbook(io.BytesIO(build_excel_report(daily, monthly)))
    for sheet_name in ("Daily", "Monthly"):
        for cell in workbook[sheet_name][1]:
            assert cell.font.bold is True


def test_header_row_is_frozen_on_both_sheets():
    daily = pd.DataFrame([{"Rep": "A", "Total Coins": 40}])
    monthly = pd.DataFrame([{"Rep": "A", "Month": "2026-07", "Total Coins": 40}])
    workbook = openpyxl.load_workbook(io.BytesIO(build_excel_report(daily, monthly)))
    for sheet_name in ("Daily", "Monthly"):
        assert workbook[sheet_name].freeze_panes == "A2"
